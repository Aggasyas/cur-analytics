# -*- coding: utf-8 -*-
"""
Парсер выгрузки тепловой карты ЦУР (Excel) -> JSON.
Фокус: настроения жителей за вчера. Решённость/закрытость НЕ анализируем.
Разрезы: направления, синт.группы, подтемы, источники, тип (проблема/предложение),
зона ответственности (администрация округа / областные-иные), критичные темы, аварии/инциденты.
"""
import sys, os, json, re
from collections import Counter, OrderedDict

import openpyxl


# ---- Критичные темы (всегда показываем, даже если 0) ----
# Каждая тема = ключ -> (заголовок, иконка, regex по направлению+группе+подтеме)
CRITICAL = OrderedDict([
    ("roads",     ("Дороги",                  "\U0001F6A7", r"дорог|ям[аы]|покрыт|асфальт|обочин|знак|разметк|ИДН|мост")),
    ("garbage",   ("Мусор и контейнеры",      "\U0001F5D1", r"контейнер|мусор|ТКО|свал|навал|отход|вывоз")),
    ("territory", ("Содержание территории",   "\U0001F333", r"содержание территор|благоустр|покос|трав|озелен|урн")),
    ("dip",       ("Детские площадки (ДИП)",  "\U0001F6DD", r"детск|игров|ДИП|площадк")),
    ("hogweed",   ("Борщевик",                "\U0001F331", r"борщевик")),
])

# Зона ответственности
ADM_PREFIX = "Администрация г. о. Солнечногорск"
NO_OWNER = {"none", "без управления", ""}

# Аварии/инциденты — срочное
INCIDENT_STATUS = {"инцидент"}
INCIDENT_SOURCE = {"инцидент"}
EMERGENCY_RE = re.compile(r"авари", re.I)


def _norm(v):
    return "" if v is None else str(v).strip()


def _is_admin(executor):
    return executor.startswith(ADM_PREFIX)


def _ownership(executor):
    e = executor.strip()
    if e.lower() in NO_OWNER:
        return "none"
    if _is_admin(e):
        return "admin"
    return "region"


def _topn(counter, n=None):
    items = counter.most_common(n)
    return [{"name": k, "count": v} for k, v in items]


def parse(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    # карта заголовков
    H = {}
    for c in range(1, ws.max_column + 1):
        name = _norm(ws.cell(1, c).value)
        if name:
            H[name] = c

    def colname(*cands):
        for cand in cands:
            for k in H:
                if k.lower() == cand.lower():
                    return k
        # частичное совпадение
        for cand in cands:
            for k in H:
                if cand.lower() in k.lower():
                    return k
        return None

    C_NAPR = colname("Направление")
    C_SYNT = colname("Синт. группа", "Синт группа")
    C_POD = colname("Подтема")
    C_STAT = colname("Статус")
    C_EXEC = colname("Исполнитель")
    C_SRC = colname("Источник")
    C_TYPE = colname("Тип сообщения - 0 проблемы, 1 - предложения", "Тип сообщения")
    C_SPAM = colname("Спам (да/нет)", "Спам")
    C_DATE = colname("Дата (первого взятия в работу)", "Дата")
    C_OMSU = colname("ОМСУ")

    rows = []
    for r in range(2, ws.max_row + 1):
        # пропускаем полностью пустые строки
        napr = _norm(ws.cell(r, H[C_NAPR]).value) if C_NAPR else ""
        synt = _norm(ws.cell(r, H[C_SYNT]).value) if C_SYNT else ""
        pod = _norm(ws.cell(r, H[C_POD]).value) if C_POD else ""
        if not (napr or synt or pod):
            continue
        rows.append({
            "napr": napr,
            "synt": synt,
            "pod": pod,
            "status": _norm(ws.cell(r, H[C_STAT]).value) if C_STAT else "",
            "exec": _norm(ws.cell(r, H[C_EXEC]).value) if C_EXEC else "",
            "src": _norm(ws.cell(r, H[C_SRC]).value) if C_SRC else "",
            "type": _norm(ws.cell(r, H[C_TYPE]).value) if C_TYPE else "",
            "spam": _norm(ws.cell(r, H[C_SPAM]).value) if C_SPAM else "",
        })

    # дата периода — из имени файла или из колонки даты
    date = _extract_date(xlsx_path, rows, ws, H, C_DATE)

    total = len(rows)

    # тип: проблемы / предложения
    problems = sum(1 for x in rows if x["type"] in ("0", "0.0"))
    proposals = sum(1 for x in rows if x["type"] in ("1", "1.0"))

    # направления
    napr_counter = Counter(x["napr"] for x in rows if x["napr"])
    synt_counter = Counter(x["synt"] for x in rows if x["synt"])
    pod_counter = Counter(x["pod"] for x in rows if x["pod"])

    # зона ответственности
    own = Counter(_ownership(x["exec"]) for x in rows)
    region_detail = Counter(
        x["exec"] for x in rows if _ownership(x["exec"]) == "region"
    )

    # источники
    src_counter = Counter(x["src"] for x in rows if x["src"])

    # критичные темы
    critical = OrderedDict()
    for key, (title, icon, pat) in CRITICAL.items():
        rx = re.compile(pat, re.I)
        cnt = sum(
            1 for x in rows
            if rx.search(" ".join([x["napr"], x["synt"], x["pod"]]))
        )
        # подтемы внутри критичной темы
        sub = Counter(
            x["pod"] for x in rows
            if x["pod"] and rx.search(" ".join([x["napr"], x["synt"], x["pod"]]))
        )
        critical[key] = {
            "title": title, "icon": icon,
            "count": cnt,
            "subtopics": _topn(sub, 6),
        }

    # аварии / инциденты (срочное)
    urgent = []
    seen = set()
    for x in rows:
        is_inc = (
            x["status"].lower() in INCIDENT_STATUS
            or x["src"].lower() in INCIDENT_SOURCE
            or EMERGENCY_RE.search(" ".join([x["synt"], x["pod"]]))
        )
        if is_inc:
            label = x["pod"] or x["synt"] or x["napr"]
            key = (x["napr"], label)
            urgent.append({"napr": x["napr"], "label": label})
            seen.add(key)
    urgent_count = len(urgent)
    urgent_by_napr = Counter(u["napr"] for u in urgent if u["napr"])

    data = {
        "meta": {
            "date": date,
            "source": "Тепловая карта ЦУР",
            "file": os.path.basename(xlsx_path),
        },
        "total": total,
        "problems": problems,
        "proposals": proposals,
        "ownership": {
            "admin": own.get("admin", 0),
            "region": own.get("region", 0),
            "none": own.get("none", 0),
            "region_detail": _topn(region_detail),
        },
        "directions": _topn(napr_counter),
        "synt_groups": _topn(synt_counter),
        "subtopics": _topn(pod_counter),
        "sources": _topn(src_counter),
        "critical": critical,
        "urgent": {
            "count": urgent_count,
            "by_direction": _topn(urgent_by_napr),
        },
    }
    return data


def _extract_date(path, rows, ws, H, C_DATE):
    """Дата периода выгрузки. Имя файла: ...02.06.2026-02.06.2026.xlsx (конечная дата периода)."""
    base = os.path.basename(path)
    # ищем DD.MM.YYYY (берём последнюю — конец периода)
    dates = re.findall(r"(\d{2})\.(\d{2})\.(\d{4})", base)
    if dates:
        d, m, y = dates[-1]
        return f"{y}-{m}-{d}"
    # из колонки "Дата (первого взятия в работу)" — берём моду
    if C_DATE and rows:
        vals = []
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, H[C_DATE]).value
            if v is None:
                continue
            s = str(v)[:10]
            mm = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
            if mm:
                vals.append(mm.group(0))
        if vals:
            return Counter(vals).most_common(1)[0][0]
    return ""


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("usage: python3 parse_cur.py <xlsx>", file=sys.stderr)
        sys.exit(1)
    d = parse(sys.argv[1])
    print(json.dumps(d, ensure_ascii=False, indent=2))
